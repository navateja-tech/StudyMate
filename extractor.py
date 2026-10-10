import os
import io
import re
from typing import Dict, Any, List, Optional, Tuple
from PIL import Image
import pymupdf as fitz
from docx import Document as DocxDocument
from pptx import Presentation

class ExtractedImage:
    """Represents an image extracted from a document or presentation."""
    def __init__(
        self,
        image_id: str,
        unit_number: int,
        unit_type: str,
        image_bytes: bytes,
        image_format: str,
        width: int,
        height: int,
        caption: Optional[str] = None
    ):
        self.image_id = image_id
        self.unit_number = unit_number
        self.unit_type = unit_type
        self.image_bytes = image_bytes
        self.image_format = (image_format or "png").lower().replace("jpeg", "jpg")
        self.width = width
        self.height = height
        self.caption = caption or f"Image from {unit_type} {unit_number}"

class ExtractedTable:
    """Represents a structured table extracted from a document or presentation."""
    def __init__(
        self,
        table_id: str,
        unit_number: int,
        unit_type: str,
        headers: List[str],
        rows: List[List[str]],
        title: Optional[str] = None
    ):
        self.table_id = table_id
        self.unit_number = unit_number
        self.unit_type = unit_type
        self.headers = headers or []
        self.rows = rows or []
        self.title = title or f"Table on {unit_type} {unit_number}"

class DocumentUnit:
    """Represents a single logical unit (e.g. Page, Slide, or Section)."""
    def __init__(
        self,
        unit_number: int,
        unit_type: str,  # "Page", "Slide", or "Section"
        text: str,
        heading: Optional[str] = None,
        images: Optional[List[ExtractedImage]] = None,
        tables: Optional[List[ExtractedTable]] = None,
        speaker_notes: Optional[str] = None
    ):
        self.unit_number = unit_number
        self.unit_type = unit_type
        self.text = (text or "").strip()
        self.heading = heading or f"{unit_type} {unit_number}"
        self.images = images if images is not None else []
        self.tables = tables if tables is not None else []
        self.speaker_notes = speaker_notes
        self.word_count = len(self.text.split()) if self.text else 0

class DocumentExtractionResult:
    """Consolidated document extraction result containing units, outline, and media."""
    def __init__(
        self,
        filename: str,
        file_type: str,
        file_size_bytes: int,
        units: List[DocumentUnit],
        headings: List[str],
        all_images: Optional[List[ExtractedImage]] = None,
        all_tables: Optional[List[ExtractedTable]] = None
    ):
        self.filename = filename
        self.file_type = file_type
        self.file_size_bytes = file_size_bytes
        self.units = units if units is not None else []
        self.headings = headings if headings is not None else []
        self.total_units = len(self.units)
        self.total_words = sum(getattr(u, 'word_count', 0) for u in self.units)

        # Aggregate all images across units if not supplied explicitly
        if all_images is not None:
            self.all_images: List[ExtractedImage] = all_images
        else:
            self.all_images: List[ExtractedImage] = []
            for u in self.units:
                if hasattr(u, "images") and u.images:
                    self.all_images.extend(u.images)

        # Aggregate all tables across units if not supplied explicitly
        if all_tables is not None:
            self.all_tables: List[ExtractedTable] = all_tables
        else:
            self.all_tables: List[ExtractedTable] = []
            for u in self.units:
                if hasattr(u, "tables") and u.tables:
                    self.all_tables.extend(u.tables)

    def __getattr__(self, name: str):
        """Defensive attribute fallback to prevent AttributeError."""
        if name == "all_images":
            imgs = []
            for u in getattr(self, "units", []):
                if hasattr(u, "images") and u.images:
                    imgs.extend(u.images)
            self.all_images = imgs
            return imgs
        elif name == "all_tables":
            tbls = []
            for u in getattr(self, "units", []):
                if hasattr(u, "tables") and u.tables:
                    tbls.extend(u.tables)
            self.all_tables = tbls
            return tbls
        raise AttributeError(f"'DocumentExtractionResult' object has no attribute '{name}'")


class DocumentExtractor:
    """Multi-format deep extractor for PDF, PPTX, DOCX, TXT, and Images."""

    @staticmethod
    def extract_from_bytes(filename: str, file_bytes: bytes) -> DocumentExtractionResult:
        ext = os.path.splitext(filename)[1].lower()
        file_size = len(file_bytes)

        if ext == ".pdf":
            return DocumentExtractor._extract_pdf(filename, file_bytes, file_size)
        elif ext in [".docx", ".doc"]:
            return DocumentExtractor._extract_docx(filename, file_bytes, file_size)
        elif ext in [".pptx", ".ppt"]:
            return DocumentExtractor._extract_pptx(filename, file_bytes, file_size)
        elif ext in [".png", ".jpg", ".jpeg", ".webp", ".bmp"]:
            return DocumentExtractor._extract_standalone_image(filename, file_bytes, file_size)
        else:
            return DocumentExtractor._extract_text(filename, file_bytes, file_size)

    @staticmethod
    def _extract_pdf(filename: str, file_bytes: bytes, file_size: int) -> DocumentExtractionResult:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        units: List[DocumentUnit] = []
        headings: List[str] = []
        seen_headings = set()
        img_counter = 1

        for page_idx in range(len(doc)):
            page = doc[page_idx]
            page_num = page_idx + 1

            # 1. Text extraction
            text = page.get_text("text").strip()

            # 2. Heading discovery
            page_heading = None
            try:
                blocks = page.get_text("blocks")
                for b in blocks:
                    block_text = b[4].strip()
                    if not block_text:
                        continue
                    first_line = block_text.split("\n")[0].strip()
                    if (re.match(r'^(chapter|unit|section|part|\d+(\.\d+)*)\b', first_line, re.IGNORECASE) or
                        (len(first_line) < 70 and (first_line.isupper() or first_line.istitle())) or
                        first_line.startswith("#")):
                        clean_h = first_line.lstrip("#").strip()
                        page_heading = clean_h
                        if clean_h not in seen_headings:
                            headings.append(clean_h)
                            seen_headings.add(clean_h)
                        break
            except Exception:
                pass

            # 3. Embedded Images extraction
            page_images: List[ExtractedImage] = []
            try:
                image_list = page.get_images(full=True)
                for img_info in image_list:
                    try:
                        xref = img_info[0]
                        base_image = doc.extract_image(xref)
                        if not base_image:
                            continue
                        img_bytes = base_image.get("image")
                        img_ext = base_image.get("ext", "png")
                        width = base_image.get("width", 0)
                        height = base_image.get("height", 0)

                        # Filter out tiny icon/spacer bullets (< 40x40)
                        if width >= 40 and height >= 40 and img_bytes:
                            page_images.append(ExtractedImage(
                                image_id=f"pdf_img_{img_counter}",
                                unit_number=page_num,
                                unit_type="Page",
                                image_bytes=img_bytes,
                                image_format=img_ext,
                                width=width,
                                height=height,
                                caption=f"Diagram/Figure {img_counter} (Page {page_num}, {width}x{height} px)"
                            ))
                            img_counter += 1
                    except Exception:
                        continue
            except Exception:
                pass

            # 4. Tables extraction
            page_tables: List[ExtractedTable] = []
            try:
                tables = page.find_tables()
                for t_idx, table in enumerate(tables):
                    table_data = table.extract()
                    if table_data and len(table_data) > 1:
                        headers = [str(col).strip() if col else f"Col {c_idx+1}" for c_idx, col in enumerate(table_data[0])]
                        rows = [[str(cell).strip() if cell else "" for cell in row] for row in table_data[1:]]
                        page_tables.append(ExtractedTable(
                            table_id=f"pdf_tbl_{page_num}_{t_idx+1}",
                            unit_number=page_num,
                            unit_type="Page",
                            headers=headers,
                            rows=rows,
                            title=f"Table {t_idx+1} on Page {page_num}"
                        ))
            except Exception:
                pass

            # Fallback label if page is visual only
            if not text and page_images:
                text = f"(Page {page_num} contains visual diagrams or images without selectable text)"
            elif not text:
                text = f"(Page {page_num} is blank or contains no selectable text)"

            units.append(DocumentUnit(
                unit_number=page_num,
                unit_type="Page",
                text=text,
                heading=page_heading or f"Page {page_num}",
                images=page_images,
                tables=page_tables
            ))

        doc.close()

        return DocumentExtractionResult(
            filename=filename,
            file_type="PDF Document",
            file_size_bytes=file_size,
            units=units,
            headings=headings
        )

    @staticmethod
    def _extract_pptx(filename: str, file_bytes: bytes, file_size: int) -> DocumentExtractionResult:
        stream = io.BytesIO(file_bytes)
        prs = Presentation(stream)
        units: List[DocumentUnit] = []
        headings: List[str] = []
        seen_headings = set()
        img_counter = 1

        for i, slide in enumerate(prs.slides):
            slide_num = i + 1
            slide_texts = []
            slide_title = None
            slide_images: List[ExtractedImage] = []
            slide_tables: List[ExtractedTable] = []

            for shape in slide.shapes:
                # Text content
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            if not slide_title and len(text) < 80:
                                slide_title = text
                            slide_texts.append(text)

                # Embedded images
                if hasattr(shape, "image"):
                    try:
                        img = shape.image
                        img_bytes = img.blob
                        img_ext = img.ext or "png"
                        with Image.open(io.BytesIO(img_bytes)) as pil_img:
                            w, h = pil_img.size

                        if w >= 40 and h >= 40:
                            slide_images.append(ExtractedImage(
                                image_id=f"pptx_img_{img_counter}",
                                unit_number=slide_num,
                                unit_type="Slide",
                                image_bytes=img_bytes,
                                image_format=img_ext,
                                width=w,
                                height=h,
                                caption=f"Slide {slide_num} Image ({w}x{h})"
                            ))
                            img_counter += 1
                    except Exception:
                        pass

                # Structured tables
                if shape.has_table:
                    try:
                        tbl = shape.table
                        t_rows = []
                        for row in tbl.rows:
                            row_cells = [cell.text.strip() for cell in row.cells]
                            t_rows.append(row_cells)
                        if t_rows:
                            headers = t_rows[0]
                            data_rows = t_rows[1:] if len(t_rows) > 1 else []
                            slide_tables.append(ExtractedTable(
                                table_id=f"pptx_tbl_{slide_num}",
                                unit_number=slide_num,
                                unit_type="Slide",
                                headers=headers,
                                rows=data_rows,
                                title=f"Slide {slide_num} Table"
                            ))
                    except Exception:
                        pass

            # Speaker Notes
            speaker_notes = None
            if slide.has_notes_slide:
                notes_frame = slide.notes_slide.notes_text_frame
                if notes_frame:
                    speaker_notes = notes_frame.text.strip()

            full_text = "\n".join(slide_texts) if slide_texts else "(Visual / Blank Slide)"
            if slide_title and slide_title not in seen_headings:
                headings.append(slide_title)
                seen_headings.add(slide_title)

            units.append(DocumentUnit(
                unit_number=slide_num,
                unit_type="Slide",
                text=full_text,
                heading=slide_title or f"Slide {slide_num}",
                images=slide_images,
                tables=slide_tables,
                speaker_notes=speaker_notes
            ))

        if not units:
            units.append(DocumentUnit(
                unit_number=1,
                unit_type="Slide",
                text="(Empty Presentation)",
                heading="Presentation"
            ))

        return DocumentExtractionResult(
            filename=filename,
            file_type="PowerPoint Presentation",
            file_size_bytes=file_size,
            units=units,
            headings=headings
        )

    @staticmethod
    def _extract_docx(filename: str, file_bytes: bytes, file_size: int) -> DocumentExtractionResult:
        stream = io.BytesIO(file_bytes)
        doc = DocxDocument(stream)
        units: List[DocumentUnit] = []
        headings: List[str] = []
        seen_headings = set()

        # Extract embedded images from Word document parts
        doc_images: List[ExtractedImage] = []
        img_counter = 1
        for rel in doc.part.related_parts.values():
            if "image" in rel.content_type:
                try:
                    img_bytes = rel.blob
                    with Image.open(io.BytesIO(img_bytes)) as pil_img:
                        w, h = pil_img.size
                        fmt = pil_img.format or "png"
                    if w >= 40 and h >= 40:
                        doc_images.append(ExtractedImage(
                            image_id=f"docx_img_{img_counter}",
                            unit_number=1,
                            unit_type="Section",
                            image_bytes=img_bytes,
                            image_format=fmt.lower(),
                            width=w,
                            height=h,
                            caption=f"Embedded Image {img_counter} ({w}x{h})"
                        ))
                        img_counter += 1
                except Exception:
                    pass

        # Extract tables from docx
        doc_tables: List[ExtractedTable] = []
        for t_idx, table in enumerate(doc.tables):
            table_rows = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                table_rows.append(row_cells)
            if table_rows:
                headers = table_rows[0]
                data_rows = table_rows[1:] if len(table_rows) > 1 else []
                doc_tables.append(ExtractedTable(
                    table_id=f"docx_tbl_{t_idx + 1}",
                    unit_number=t_idx + 1,
                    unit_type="Section",
                    headers=headers,
                    rows=data_rows,
                    title=f"Document Table {t_idx + 1}"
                ))

        # Extract sections and paragraphs
        current_heading = "Overview"
        current_paras = []
        section_idx = 1

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            is_heading = p.style.name.startswith("Heading") or re.match(r'^(chapter|unit|section|\d+\.)\b', text, re.IGNORECASE)
            if is_heading and len(text) < 100:
                if current_paras:
                    units.append(DocumentUnit(
                        unit_number=section_idx,
                        unit_type="Section",
                        text="\n\n".join(current_paras),
                        heading=current_heading
                    ))
                    section_idx += 1
                    current_paras = []
                current_heading = text
                if text not in seen_headings:
                    headings.append(text)
                    seen_headings.add(text)
            else:
                current_paras.append(text)

        if current_paras:
            units.append(DocumentUnit(
                unit_number=section_idx,
                unit_type="Section",
                text="\n\n".join(current_paras),
                heading=current_heading
            ))

        if not units:
            units.append(DocumentUnit(
                unit_number=1,
                unit_type="Section",
                text="(No text extracted from document)",
                heading="Overview"
            ))

        # Attach images and tables to the first unit if available
        if units:
            units[0].images = doc_images
            units[0].tables = doc_tables

        return DocumentExtractionResult(
            filename=filename,
            file_type="Word DOCX Document",
            file_size_bytes=file_size,
            units=units,
            headings=headings
        )

    @staticmethod
    def _extract_standalone_image(filename: str, file_bytes: bytes, file_size: int) -> DocumentExtractionResult:
        with Image.open(io.BytesIO(file_bytes)) as pil_img:
            w, h = pil_img.size
            fmt = pil_img.format or "PNG"
            color_mode = pil_img.mode

        img_obj = ExtractedImage(
            image_id="uploaded_img_1",
            unit_number=1,
            unit_type="Image",
            image_bytes=file_bytes,
            image_format=fmt.lower(),
            width=w,
            height=h,
            caption=f"{filename} ({w}x{h}, {color_mode})"
        )

        unit = DocumentUnit(
            unit_number=1,
            unit_type="Image",
            text=f"Uploaded Image: {filename}\nResolution: {w}x{h} px\nFormat: {fmt}\nColor Mode: {color_mode}",
            heading=filename,
            images=[img_obj]
        )

        return DocumentExtractionResult(
            filename=filename,
            file_type=f"Image ({fmt})",
            file_size_bytes=file_size,
            units=[unit],
            headings=[filename]
        )

    @staticmethod
    def _extract_text(filename: str, file_bytes: bytes, file_size: int) -> DocumentExtractionResult:
        content = file_bytes.decode("utf-8", errors="replace")
        paragraphs = content.split("\n\n")
        units: List[DocumentUnit] = []
        headings: List[str] = []
        seen_headings = set()

        current_heading = "Introduction"
        current_paras = []
        section_idx = 1

        for p in paragraphs:
            text = p.strip()
            if not text:
                continue

            is_heading = text.startswith("#") or (re.match(r'^(chapter|unit|section|\d+\.)\b', text, re.IGNORECASE) and len(text) < 90)
            if is_heading:
                clean_h = text.lstrip("#").strip()
                if current_paras:
                    units.append(DocumentUnit(
                        unit_number=section_idx,
                        unit_type="Section",
                        text="\n\n".join(current_paras),
                        heading=current_heading
                    ))
                    section_idx += 1
                    current_paras = []
                current_heading = clean_h
                if clean_h not in seen_headings:
                    headings.append(clean_h)
                    seen_headings.add(clean_h)
            else:
                current_paras.append(text)

        if current_paras:
            units.append(DocumentUnit(
                unit_number=section_idx,
                unit_type="Section",
                text="\n\n".join(current_paras),
                heading=current_heading
            ))

        if not units:
            units.append(DocumentUnit(
                unit_number=1,
                unit_type="Section",
                text=content or "(Empty file)",
                heading="Content"
            ))

        return DocumentExtractionResult(
            filename=filename,
            file_type="Plain Text / Markdown",
            file_size_bytes=file_size,
            units=units,
            headings=headings
        )
